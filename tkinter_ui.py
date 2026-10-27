import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from customer import Customer
from menu import Menu
from order import Order
import pickle
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime


def run_app():
    root = tk.Tk()

    root.title("سیستم مدیریت رستوران")
    root.geometry("900x600")
    root.minsize(700, 450)

    title = ttk.Label(
        root,
        text="سیستم مدیریت رستوران",
        font=("Tahoma", 20, "bold")
    )

    title.pack(pady=30)

    notebook = ttk.Notebook(root)
    notebook.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=10
    )

    customer_tab = ttk.Frame(notebook)
    menu_tab = ttk.Frame(notebook)
    order_tab = ttk.Frame(notebook)
    admin_tab = ttk.Frame(notebook)
    notebook.add(admin_tab, text="مدیریت مشتریان")

    orders_admin_tab = ttk.Frame(notebook)

    notebook.add(
        orders_admin_tab,
        text="مدیریت سفارش‌ها"
    )

    orders_search_frame = ttk.Frame(orders_admin_tab)
    orders_search_frame.pack(fill="x", padx=20, pady=10)

    ttk.Label(
        orders_search_frame,
        text="جست‌وجوی سفارش:"
    ).pack(side="right", padx=10)

    orders_search_var = tk.StringVar()

    orders_search_entry = ttk.Entry(
        orders_search_frame,
        textvariable=orders_search_var,
        width=35,
        justify="right"
    )

    orders_search_entry.pack(side="right", padx=5)

    orders_admin_frame = ttk.LabelFrame(
        orders_admin_tab,
        text="فهرست سفارش‌های ثبت‌شده"
    )

    orders_admin_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=15
    )

    orders_table_frame = ttk.Frame(orders_admin_frame)
    orders_table_frame.pack(fill="x", padx=10, pady=5)

    orders_admin_table = ttk.Treeview(
        orders_table_frame,
        columns=("id", "customer", "total"),
        show="headings",
        height=6
    )

    orders_admin_table.heading(
        "id", text="شماره سفارش"
    )

    orders_admin_table.heading(
        "customer", text="نام مشتری"
    )

    orders_admin_table.heading(
        "total", text="مبلغ کل (تومان)"
    )

    for column in ("id", "customer", "total"):
        orders_admin_table.column(
            column,
            anchor="center",
            width=180
        )

    orders_scrollbar = ttk.Scrollbar(
        orders_table_frame,
        orient="vertical",
        command=orders_admin_table.yview
    )

    orders_admin_table.configure(
        yscrollcommand=orders_scrollbar.set
    )

    orders_admin_table.pack(
        side="left",
        fill="both",
        expand=True
    )

    orders_scrollbar.pack(
        side="right",
        fill="y"
    )

    sales_summary_frame = ttk.LabelFrame(
        orders_admin_tab,
        text="خلاصه آمار فروش"
    )

    sales_summary_frame.pack(
        fill="x",
        padx=20,
        pady=5
    )

    orders_count_var = tk.StringVar(value="تعداد سفارش‌ها: 0")
    total_sales_var = tk.StringVar(value="مجموع فروش: 0 تومان")
    average_sales_var = tk.StringVar(value="میانگین سفارش: 0 تومان")

    ttk.Label(
        sales_summary_frame,
        textvariable=orders_count_var
    ).pack(side="right", padx=20, pady=10)

    ttk.Label(
        sales_summary_frame,
        textvariable=total_sales_var
    ).pack(side="right", padx=20, pady=10)

    ttk.Label(
        sales_summary_frame,
        textvariable=average_sales_var
    ).pack(side="right", padx=20, pady=10)

    def sort_orders(column, reverse=False):
        items = []

        for item_id in orders_admin_table.get_children():
            value = orders_admin_table.set(item_id, column)

            if column in ("id", "total"):
                value = int(str(value).replace(",", ""))

            items.append((value, item_id))

        items.sort(key=lambda item: item[0], reverse=reverse)

        for index, (_, item_id) in enumerate(items):
            orders_admin_table.move(item_id, "", index)

        orders_admin_table.heading(
            column,
            command=lambda: sort_orders(column, not reverse)
        )

    orders_admin_table.heading(
        "id",
        command=lambda: sort_orders("id")
    )

    orders_admin_table.heading(
        "customer",
        command=lambda: sort_orders("customer")
    )

    orders_admin_table.heading(
        "total",
        command=lambda: sort_orders("total")
    )

    def refresh_orders_admin():
        try:
            orders = {}

            orders_file = "data/orders.pkl"

            if os.path.exists(orders_file):
                with open(orders_file, "rb") as file:
                    orders = pickle.load(file)

            if not isinstance(orders, dict):
                raise ValueError(
                    "ساختار اطلاعات سفارش‌ها معتبر نیست."
                )

            rows = []
            search_text = orders_search_var.get().strip().lower().replace("ي", "ی").replace("ك", "ک")

            for order in orders.values():
                customer = order.customer

                customer_name = (
                    f"{customer.first_name} "
                    f"{customer.last_name}"
                )

                order_id_text = str(order.order_id)

                if search_text:
                    if (
                        search_text not in order_id_text
                        and search_text not in customer_name.lower().replace("ي", "ی").replace("ك", "ک")
                    ):
                        continue

                rows.append((
                    order.order_id,
                    customer_name,
                    f"{order.calculate_total():,}"
                ))

            for item in orders_admin_table.get_children():
                orders_admin_table.delete(item)

            for row in rows:
                orders_admin_table.insert(
                    "",
                    "end",
                    values=row
                )

            order_count = len(rows)

            total_sales = sum(
                int(str(row[2]).replace(",", ""))
                for row in rows
            )

            average_sales = (
                total_sales / order_count
                if order_count > 0 else 0
            )

            orders_count_var.set(
                f"تعداد سفارش‌ها: {order_count}"
            )

            total_sales_var.set(
                f"مجموع فروش: {total_sales:,.0f} تومان"
            )

            average_sales_var.set(
                f"میانگین سفارش: {average_sales:,.0f} تومان"
            )

        except (
            OSError,
            ValueError,
            TypeError,
            EOFError,
            AttributeError,
            pickle.UnpicklingError
        ) as error:
            messagebox.showerror(
                "خطای بازیابی سفارش‌ها",
                str(error)
            )

    def on_orders_search_change(_name, _index, _mode):
        refresh_orders_admin()
        return ""

    orders_search_var.trace_add(
        "write",
        on_orders_search_change
    )

    orders_refresh_button = ttk.Button(
        orders_admin_tab,
        text="نمایش / به‌روزرسانی سفارش‌ها",
        command=refresh_orders_admin
    )

    orders_refresh_button.pack(pady=8)

    def show_order_details():
        selected = orders_admin_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب سفارش",
                "لطفاً ابتدا یک سفارش انتخاب کنید."
            )
            return

        order_id = int(
            orders_admin_table.item(selected[0], "values")[0]
        )

        try:
            with open("data/orders.pkl", "rb") as file:
                orders = pickle.load(file)

            order = orders.get(order_id)

            if order is None:
                messagebox.showerror(
                    "خطا",
                    "سفارش موردنظر پیدا نشد."
                )
                return

            details_window = tk.Toplevel(root)
            details_window.title(f"جزئیات سفارش {order_id}")
            details_window.geometry("550x450")

            customer = order.customer
            customer_name = (
                f"{customer.first_name} {customer.last_name}"
            )

            ttk.Label(
                details_window,
                text=f"شماره سفارش: {order_id}",
                font=("Tahoma", 13, "bold")
            ).pack(pady=10)

            ttk.Label(
                details_window,
                text=f"نام مشتری: {customer_name}"
            ).pack(pady=5)

            details_table = ttk.Treeview(
                details_window,
                columns=("food", "quantity", "price", "total"),
                show="headings",
                height=8
            )

            for column, title in [
                ("food", "نام غذا"),
                ("quantity", "تعداد"),
                ("price", "قیمت واحد"),
                ("total", "مبلغ")
            ]:
                details_table.heading(column, text=title)
                details_table.column(
                    column, width=120, anchor="center"
                )

            details_table.pack(
                fill="both", expand=True, padx=15, pady=10
            )

            for food_name, item in order.get_items().items():
                quantity = item["quantity"]
                price = item["price"]

                details_table.insert(
                    "", "end",
                    values=(
                        food_name,
                        quantity,
                        f"{price:,}",
                        f"{price * quantity:,}"
                    )
                )

            ttk.Label(
                details_window,
                text=f"مبلغ کل: {order.calculate_total():,} تومان",
                font=("Tahoma", 12, "bold")
            ).pack(pady=10)

        except (OSError, KeyError, TypeError, ValueError, AttributeError, pickle.UnpicklingError) as error:
            messagebox.showerror(
                "خطا در نمایش سفارش",
                str(error)
            )

    def export_orders_to_excel():
        try:
            file_path = filedialog.asksaveasfilename(
                title="ذخیره گزارش سفارش‌ها",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile="restaurant_orders.xlsx"
            )

            if not file_path:
                return

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "سفارش‌ها"
            sheet.sheet_view.rightToLeft = True

            sheet.merge_cells("A1:C1")
            sheet["A1"] = "گزارش فروش رستوران"

            sheet["A1"].font = Font(
                name="Tahoma",
                size=16,
                bold=True,
                color="FFFFFF"
            )

            sheet["A1"].fill = PatternFill(
                fill_type="solid",
                fgColor="243B55"
            )

            sheet["A1"].alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            sheet.row_dimensions[1].height = 35

            sheet.merge_cells("A2:C2")
            sheet["A2"] = (
                "تاریخ تهیه گزارش: "
                + datetime.now().strftime("%Y/%m/%d")
            )

            sheet["A2"].alignment = Alignment(
                horizontal="center"
            )

            sheet.append([])
            sheet.append([])
            sheet.append([])
            sheet.append([
                "شماره سفارش",
                "نام مشتری",
                "مبلغ کل (تومان)"
            ])

            for item_id in orders_admin_table.get_children():
                values = orders_admin_table.item(
                    item_id, "values"
                )

                sheet.append([
                    int(values[0]),
                    values[1],
                    int(str(values[2]).replace(",", ""))
                ])

            for cell in sheet[6]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="243B55"
                )
                cell.alignment = Alignment(
                    horizontal="center"
                )

            sheet.column_dimensions["A"].width = 20
            sheet.column_dimensions["B"].width = 30
            sheet.column_dimensions["C"].width = 25

            for row in sheet.iter_rows(min_row=7):
                for cell in row:
                    cell.alignment = Alignment(
                        horizontal="center"
                    )
                row[2].number_format = "#,##0"

            order_count = len(orders_admin_table.get_children())

            total_sales = sum(
                sheet.cell(row=i, column=3).value
                for i in range(5, sheet.max_row + 1)
                if isinstance(
                    sheet.cell(row=i, column=3).value,
                    (int, float)
                )
            )

            average_sales = (
                total_sales / order_count
                if order_count > 0 else 0
            )

            summary_row = sheet.max_row + 2

            sheet.cell(
                row=summary_row,
                column=2,
                value="تعداد سفارش‌ها"
            )
            sheet.cell(
                row=summary_row,
                column=3,
                value=order_count
            )

            sheet.cell(
                row=summary_row + 1,
                column=2,
                value="مجموع فروش (تومان)"
            )
            sheet.cell(
                row=summary_row + 1,
                column=3,
                value=total_sales
            )

            sheet.cell(
                row=summary_row + 2,
                column=2,
                value="میانگین سفارش (تومان)"
            )
            sheet.cell(
                row=summary_row + 2,
                column=3,
                value=average_sales
            )

            for row in range(summary_row, summary_row + 3):
                sheet.cell(row=row, column=2).font = Font(bold=True)
                sheet.cell(row=row, column=3).font = Font(bold=True)
                sheet.cell(
                    row=row,
                    column=3
                ).number_format = "#,##0"

            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.orientation = "portrait"
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1

            sheet.print_options.horizontalCentered = True
            sheet.print_area = f"A1:C{sheet.max_row}"

            sheet.page_margins.left = 0.3
            sheet.page_margins.right = 0.3
            sheet.page_margins.top = 0.5
            sheet.page_margins.bottom = 0.5

            sheet.row_dimensions[6].height = 30

            for row in range(7, 15):
                sheet.row_dimensions[row].height = 28

                for cell in sheet[row][:3]:
                    cell.font = Font(
                        name="Tahoma",
                        size=13
                    )
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center"
                    )

            for row in range(16, 19):
                sheet.row_dimensions[row].height = 30

                for cell in sheet[row][1:3]:
                    cell.font = Font(
                        name="Tahoma",
                        size=12,
                        bold=True
                    )

            workbook.save(file_path)

            messagebox.showinfo(
                "خروجی اکسل",
                "فایل اکسل با موفقیت ذخیره شد."
            )

        except (OSError, ValueError, IndexError) as error:
            messagebox.showerror(
                "خطا در خروجی اکسل",
                str(error)
            )

    details_button = ttk.Button(
        orders_admin_tab,
        text="مشاهده جزئیات سفارش",
        command=show_order_details
    )

    details_button.pack(pady=5)
    export_excel_button = ttk.Button(
        orders_admin_tab,
        text="خروجی اکسل سفارش‌ها",
        command=export_orders_to_excel
    )

    export_excel_button.pack(pady=5)

    refresh_orders_admin()

    customers_frame = ttk.LabelFrame(
        admin_tab,
        text="فهرست مشتریان"
    )
    customers_frame.pack(
        fill="both",
        expand=False,
        padx=20,
        pady=15
    )

    customers_table_frame = ttk.Frame(customers_frame)
    customers_table_frame.pack(fill="x", padx=10, pady=5)

    customers_table = ttk.Treeview(
        customers_table_frame,
        columns=("id", "first_name", "last_name", "phone"),
        show="headings",
        height=8
    )

    customers_table.heading("id", text="شناسه")
    customers_table.heading("first_name", text="نام")
    customers_table.heading("last_name", text="نام خانوادگی")
    customers_table.heading("phone", text="شماره تماس")

    customers_table.column("id", width=100, anchor="center")
    customers_table.column("first_name", width=150, anchor="center")
    customers_table.column("last_name", width=150, anchor="center")
    customers_table.column("phone", width=180, anchor="center")

    customers_scrollbar = ttk.Scrollbar(
        customers_table_frame,
        orient="vertical",
        command=customers_table.yview
    )

    customers_table.configure(
        yscrollcommand=customers_scrollbar.set
    )

    customers_table.pack(
        side="left",
        fill="both",
        expand=True
    )

    customers_scrollbar.pack(
        side="right",
        fill="y"
    )

    def refresh_customers():
        for item in customers_table.get_children():
            customers_table.delete(item)

        try:
            customers = Customer.load_all_customers()

            for customer in customers.values():
                customers_table.insert(
                    "",
                    "end",
                    values=(
                        customer.customer_id,
                        customer.first_name,
                        customer.last_name,
                        customer.phone_number
                    )
                )

        except (OSError, ValueError, EOFError, pickle.UnpicklingError) as error:
            messagebox.showerror(
                "خطا",
                str(error)
            )
    refresh_button = ttk.Button(
        admin_tab,
        text="نمایش / به‌روزرسانی مشتریان",
        command=refresh_customers
    )
    refresh_button.pack(pady=5)

    search_frame = ttk.Frame(admin_tab)
    search_frame.pack(fill="x", padx=20, pady=10)

    ttk.Label(
        search_frame,
        text="جست‌وجوی شناسه یا شماره تماس:"
    ).pack(side="right", padx=5)

    search_entry = ttk.Entry(search_frame, width=25)
    search_entry.pack(side="right", padx=5)

    def search_customer():
        query = search_entry.get().strip()

        if not query:
            refresh_customers()
            return

        try:
            customers = Customer.load_all_customers()

            for item in customers_table.get_children():
                customers_table.delete(item)

            for customer in customers.values():
                if (
                    query == str(customer.customer_id)
                    or query in customer.phone_number
                ):
                    customers_table.insert(
                        "",
                        "end",
                        values=(
                            customer.customer_id,
                            customer.first_name,
                            customer.last_name,
                            customer.phone_number
                        )
                    )

        except (OSError, ValueError, EOFError, pickle.UnpicklingError) as error:
            messagebox.showerror("خطای جست‌وجو", str(error))

    search_button = ttk.Button(
        search_frame,
        text="جست‌وجو",
        command=search_customer
    )
    search_button.pack(side="right", padx=5)

    refresh_customers()

    notebook.add(customer_tab, text="اطلاعات مشتری")
    notebook.add(menu_tab, text="منوی رستوران")
    notebook.add(order_tab, text="ثبت سفارش و فاکتور")

    # فرم اطلاعات مشتری
    customer_frame = ttk.LabelFrame(
        customer_tab,
        text="ثبت اطلاعات مشتری",
        padding=20
    )
    customer_frame.pack(padx=30, pady=25, fill="x")

    fields = [
        ("شناسه مشتری", "customer_id"),
        ("نام", "first_name"),
        ("نام خانوادگی", "last_name"),
        ("شماره تماس", "phone_number"),
        ("آدرس", "address")
    ]

    customer_entries = {}

    for row_number, (label_text, field_name) in enumerate(fields):
        ttk.Label(
            customer_frame,
            text=label_text
        ).grid(row=row_number, column=1, padx=10, pady=8, sticky="e")

        entry = ttk.Entry(customer_frame, width=35, justify="right")
        entry.grid(row=row_number, column=0, padx=10, pady=8)

        customer_entries[field_name] = entry

    ttk.Label(
        customer_frame,
        text="جنسیت"
    ).grid(row=5, column=1, padx=10, pady=8, sticky="e")

    gender_combo = ttk.Combobox(
        customer_frame,
        values=["", "مرد", "زن", "سایر"],
        state="readonly",
        width=32,
        justify="right"
    )
    gender_combo.grid(row=5, column=0, padx=10, pady=8)
    gender_combo.current(0)

    # نگهداری مشتری ثبت ‌شده
    current_customer: Customer | None = None

    def select_existing_customer():
        nonlocal current_customer

        try:
            customer_id = int(
                customer_entries["customer_id"].get().strip()
            )

            customer = Customer.load_customer(customer_id)
            current_customer = customer

            messagebox.showinfo(
                "انتخاب مشتری",
                "مشتری با موفقیت برای سفارش انتخاب شد."
            )

        except (ValueError, FileNotFoundError) as error:
            messagebox.showerror(
                "خطا در انتخاب مشتری",
                str(error)
            )

    def register_customer():
        nonlocal current_customer

        try:
            customer_id = int(
                customer_entries["customer_id"].get().strip()
            )

            current_customer = Customer(
                customer_id=customer_id,
                first_name=customer_entries["first_name"].get(),
                last_name=customer_entries["last_name"].get(),
                phone_number=customer_entries["phone_number"].get(),
                address=customer_entries["address"].get(),
                gender=gender_combo.get()
            )
            current_customer.save_customer()

            messagebox.showinfo(
                "ثبت موفق",
                "اطلاعات مشتری با موفقیت ثبت شد."
            )

        except ValueError as error:
            messagebox.showerror(
                "خطای اطلاعات مشتری",
                str(error)
            )
    register_button = ttk.Button(
        customer_frame,
        text="ثبت اطلاعات مشتری",
        command=register_customer
    )

    register_button.grid(
        row=6,
        column=0,
        columnspan=2,
        pady=20
    )

    select_customer_button = ttk.Button(
        customer_tab,
        text="انتخاب مشتری موجود",
        command=select_existing_customer
    )

    select_customer_button.pack(pady=5)

    restaurant_menu = Menu()

    try:
        restaurant_menu = Menu.load_menu()
    except FileNotFoundError:
        restaurant_menu = Menu()

        default_foods = {
            "پیتزا": 250000,
            "برگر": 180000,
            "پاستا": 220000,
            "سالاد سزار": 150000,
            "نوشابه": 30000
        }

        for food_name, price in default_foods.items():
            restaurant_menu.add_item(food_name, price)

    menu_frame = ttk.LabelFrame(
        menu_tab,
        text="لیست غذاهای رستوران",
        padding=20
    )
    menu_frame.pack(fill="both", expand=False, padx=20, pady=20)

    # قاب جدول و زبانه پیمایش
    menu_table_frame = ttk.Frame(menu_frame)
    menu_table_frame.pack(
        fill="x",
        padx=10,
        pady=10
    )

    menu_table = ttk.Treeview(
        menu_frame,
        columns=("name", "price"),
        show="headings",
        height=5
    )

    menu_table.heading("name", text="نام غذا")
    menu_table.heading("price", text="قیمت (تومان)")

    menu_table.column("name", width=250, anchor="center")
    menu_table.column("price", width=180, anchor="center")

    # زبانه عمودی
    menu_scrollbar = ttk.Scrollbar(
        menu_table_frame,
        orient="vertical",
        command=menu_table.yview
    )

    menu_table.configure(
        yscrollcommand=menu_scrollbar.set
    )

    menu_table.pack(
        in_=menu_table_frame,
        side="left",
        fill="x",
        expand=True
    )

    menu_scrollbar.pack(
        side="right",
        fill="y"
    )

    for name, price in restaurant_menu.get_items().items():
        menu_table.insert(
            "",
            "end",
            values=(name, f"{price:,}")
        )

    def refresh_menu_table():
        for item in menu_table.get_children():
            menu_table.delete(item)

        for name, price in restaurant_menu.get_items().items():
            menu_table.insert(
                "",
                "end",
                values=(name, f"{price:,}")
            )

    def add_new_food():
        name = food_name_entry.get().strip()
        price_text = food_price_entry.get().strip()

        if not name or not price_text:
            messagebox.showwarning(
                "اطلاعات ناقص",
                "نام غذا و قیمت را وارد کنید."
            )
            return

        try:
            price = int(price_text)

            if price <= 0:
                raise ValueError("قیمت باید بیشتر از صفر باشد.")

            if name in restaurant_menu.get_items():
                raise ValueError("این غذا قبلاً در منو وجود دارد.")

            restaurant_menu.add_item(name, price)
            restaurant_menu.save_menu()

            messagebox.showinfo(
                "ثبت موفق",
                "غذای جدید با موفقیت به منو اضافه شد."
            )

            food_name_entry.delete(0, "end")
            food_price_entry.delete(0, "end")

            refresh_menu_table()
            refresh_food_combobox()

        except (ValueError, TypeError, OSError) as error:
            messagebox.showerror(
                "خطای افزودن غذا",
                str(error)
            )

    food_form = ttk.LabelFrame(
        menu_tab,
        text="افزودن غذای جدید"
    )
    food_form.pack(
        fill="x",
        padx=20,
        pady=2
    )

    ttk.Label(
        food_form,
        text="نام غذا:"
    ).grid(row=0, column=1, padx=10, pady=10)

    food_name_entry = ttk.Entry(food_form, width=30)
    food_name_entry.grid(row=0, column=0, padx=10, pady=10)

    ttk.Label(
        food_form,
        text="قیمت (تومان):"
    ).grid(row=1, column=1, padx=10, pady=10)

    food_price_entry = ttk.Entry(food_form, width=30)
    food_price_entry.grid(row=1, column=0, padx=10, pady=10)

    add_food_button = ttk.Button(
        food_form,
        text="افزودن غذا به منو",
        command=add_new_food
    )
    add_food_button.grid(
        row=2,
        column=0,
        columnspan=2,
        pady=5
    )

    def edit_food_price():
        selected = menu_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب غذا",
                "لطفاً ابتدا یک غذا از جدول انتخاب کنید."
            )
            return

        values = menu_table.item(selected[0], "values")
        food_name = values[0]

        price_window = tk.Toplevel(root)
        price_window.title("ویرایش قیمت غذا")
        price_window.geometry("400x220")
        price_window.transient(root)
        price_window.grab_set()

        ttk.Label(
            price_window,
            text=f"قیمت جدید برای {food_name}:"
        ).pack(pady=15)

        price_entry = ttk.Entry(price_window, width=25)
        price_entry.pack(pady=5)

        current_price = restaurant_menu.get_items()[food_name]
        price_entry.insert(0, str(current_price))

        def save_new_price():
            try:
                new_price = int(price_entry.get().strip())

                old_price = restaurant_menu.get_items()[food_name]

                restaurant_menu.update_price(food_name, new_price)

                try:
                    restaurant_menu.save_menu()
                except OSError:
                    restaurant_menu.update_price(food_name, old_price)
                    raise

                refresh_menu_table()
                refresh_food_combobox()

                messagebox.showinfo(
                    "ویرایش موفق",
                    "قیمت غذا با موفقیت تغییر کرد.",
                    parent=price_window
                )

                price_window.destroy()

            except (ValueError, TypeError, OSError) as error:
                messagebox.showerror(
                    "خطای ویرایش قیمت",
                    str(error),
                    parent=price_window
                )

        save_button = ttk.Button(
            price_window,
            text="ذخیره قیمت جدید",
            command=save_new_price
        )

        save_button.pack(pady=15)

    edit_price_button = ttk.Button(
        food_form,
        text="ویرایش قیمت غذای انتخاب‌شده",
        command=edit_food_price
    )

    edit_price_button.grid(
        row=3,
        column=0,
        columnspan=2,
        pady=5
    )

    def delete_selected_food():
        selected = menu_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب غذا",
                "لطفاً ابتدا یک غذا از جدول انتخاب کنید."
            )
            return

        values = menu_table.item(selected[0], "values")
        food_name = values[0]

        confirmed = messagebox.askyesno(
            "تأیید حذف غذا",
            f"آیا از حذف «{food_name}» از منو مطمئن هستید؟"
        )

        if not confirmed:
            return

        try:
            old_items = restaurant_menu.get_items().copy()

            restaurant_menu.remove_item(food_name)

            try:
                restaurant_menu.save_menu()
            except OSError:
                restaurant_menu._items = old_items
                raise

            refresh_menu_table()
            refresh_food_combobox()

            messagebox.showinfo(
                "حذف موفق",
                f"غذای «{food_name}» با موفقیت حذف شد."
            )

        except (ValueError, TypeError, OSError) as error:
            messagebox.showerror(
                "خطای حذف غذا",
                str(error)
            )

    delete_food_button = ttk.Button(
        food_form,
        text="حذف غذای انتخاب‌شده",
        command=delete_selected_food
    )

    delete_food_button.grid(
        row=4,
        column=0,
        columnspan=2,
        pady=5
    )

    order_frame = ttk.LabelFrame(
        order_tab,
        text="ثبت سفارش جدید",
        padding=20
    )
    order_frame.pack(fill="x", padx=20, pady=20)

    ttk.Label(
        order_frame,
        text="انتخاب غذا:"
    ).grid(row=0, column=2, padx=10, pady=10)

    food_combo = ttk.Combobox(
        order_frame,
        values=list(restaurant_menu.get_items().keys()),
        state="readonly",
        width=25,
        justify="right"
    )
    food_combo.grid(row=0, column=1, padx=10, pady=10)

    def refresh_food_combobox():
        food_names = list(restaurant_menu.get_items().keys())

        food_combo["values"] = food_names

        current_food = food_combo.get()

        if current_food not in food_names:
            if food_names:
                food_combo.set(food_names[0])
            else:
                food_combo.set("")

    refresh_food_combobox()

    ttk.Label(
        order_frame,
        text="تعداد:"
    ).grid(row=1, column=2, padx=10, pady=10)

    quantity_spin = ttk.Spinbox(
        order_frame,
        from_=1,
        to=100,
        width=10
    )
    quantity_spin.grid(row=1, column=1, padx=10, pady=10)
    quantity_spin.set(1)
    current_order = None
    def add_food_to_order():
        nonlocal current_order

        try:
            if current_customer is None:
                messagebox.showwarning(
                    "مشتری ثبت نشده",
                    "ابتدا اطلاعات مشتری را ثبت کنید."
                )
                return

            food_name = food_combo.get().strip()
            quantity = int(quantity_spin.get())

            if not food_name:
                raise ValueError("ابتدا یک غذا انتخاب کنید.")
            if current_order is None:
                new_order_id = Order.get_next_order_id()

                customer = current_customer

                if not isinstance(customer, Customer):
                    raise ValueError(
                        "ابتدا اطلاعات مشتری را ثبت کنید."
                    )

                current_order = Order(
                    order_id=new_order_id,
                    customer=customer,
                    menu=restaurant_menu
                )
                print("شماره سفارش جدید:", current_order.order_id)

            current_order.add_item(food_name, quantity)
            update_order_table()

            messagebox.showinfo(
                "ثبت موفق",
                "غذا با موفقیت به سفارش اضافه شد."
            )

        except (ValueError, TypeError) as error:
            messagebox.showerror(
                "خطای سفارش",
                str(error)
            )
    add_food_button = ttk.Button(
        order_frame,
        text="افزودن به سفارش",
        command=add_food_to_order
    )

    add_food_button.grid(
        row=2,
        column=0,
        columnspan=3,
        pady=20
    )

    order_table_frame = ttk.Frame(order_tab)
    order_table_frame.pack(fill="x", padx=10, pady=2)

    order_table = ttk.Treeview(
        order_table_frame,
        columns=("name", "quantity", "price", "subtotal"),
        show="headings",
        height=4
    )

    order_table.heading("name", text="نام غذا")
    order_table.heading("quantity", text="تعداد")
    order_table.heading("price", text="قیمت واحد")
    order_table.heading("subtotal", text="مبلغ")

    for column in ("name", "quantity", "price", "subtotal"):
        order_table.column(column, anchor="center", width=150)

    order_scrollbar = ttk.Scrollbar(
        order_table_frame,
        orient="vertical",
        command=order_table.yview
    )

    order_table.configure(
        yscrollcommand=order_scrollbar.set
    )

    order_table.pack(
        side="left",
        fill="both",
        expand=True
    )

    order_scrollbar.pack(
        side="right",
        fill="y"
    )

    total_label = ttk.Label(
        order_tab,
        text="مبلغ کل: 0 تومان",
        font=("Tahoma", 13, "bold")
    )
    total_label.pack(pady=2)

    def update_order_table():
        for row in order_table.get_children():
            order_table.delete(row)

        order = current_order

        if not isinstance(order, Order):
            total_label.config(text="مبلغ کل: 0 تومان")
            return

        for name, details in order.get_items().items():
            quantity = details["quantity"]
            price = details["price"]

            order_table.insert(
                "",
                "end",
                values=(
                    name,
                    quantity,
                    f"{price:,}",
                    f"{quantity * price:,}"
                )
            )

        if current_order is None:
            total_label.config(text="مبلغ کل: 0 تومان")
            return

        total_label.config(
            text=f"مبلغ کل: {order.calculate_total():,} تومان"
        )

    def remove_selected_food():
        order = current_order
        if not isinstance(order, Order):
            messagebox.showwarning(
                "سفارش خالی",
                "هنوز سفارشی ثبت نشده است."
            )
            return

        selected = order_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب غذا",
                "ابتدا یک غذا از جدول انتخاب کنید."
            )
            return

        food_name = order_table.item(selected[0], "values")[0]

        try:
            order.remove_item(food_name)
            update_order_table()

            messagebox.showinfo(
                "حذف موفق",
                "غذا از سفارش حذف شد."
            )

        except ValueError as error:
            messagebox.showerror(
                "خطا",
                str(error)
            )
    remove_button = ttk.Button(
        order_tab,
        text="حذف غذای انتخاب ‌شده",
        command=remove_selected_food
    )

    remove_button.pack(
        side="top",
        pady=5,
        before=order_table
    )

    def finalize_order():
        nonlocal current_order

        if current_order is None:
            messagebox.showwarning(
                "سفارش خالی",
                "ابتدا غذا به سفارش اضافه کنید."
            )
            return

        order = current_order
        if not isinstance(order, Order):
            return

        if not order.get_items():
            messagebox.showwarning(
                "سفارش خالی",
                "هیچ غذایی در سفارش وجود ندارد."
            )
            return

        try:
            order.save_order()

            saved_id = order.order_id

            messagebox.showinfo(
                "ثبت نهایی موفق",
                f"سفارش شماره {saved_id} با موفقیت ذخیره شد."
            )
            show_invoice(order)

            current_order = None
            update_order_table()

        except (ValueError, TypeError, OSError) as error:
            messagebox.showerror(
                "خطا در ثبت سفارش",
                str(error)
            )

    def show_invoice(order):
        invoice_window = tk.Toplevel(root)
        invoice_window.transient(root)
        invoice_window.grab_set()
        invoice_window.focus_set()
        invoice_window.title("فاکتور سفارش")
        invoice_window.geometry("650x500")

        invoice_text = tk.Text(
            invoice_window,
            font=("Tahoma", 12),
            wrap="word",
            padx=20,
            pady=20
        )

        lines = [
            "فاکتور رستوران",
            "=" * 40,
            f"شماره سفارش: {order.order_id}",
            f"تاریخ و ساعت سفارش: {order.get_order_datetime()}",
            f"نام مشتری: {order.customer.first_name} "
            f"شماره تماس: {order.customer.phone_number}",
            f"آدرس: {order.customer.address}",
            f"{order.customer.last_name}",
            "",
            "اقلام سفارش:",
            "-" * 40
        ]

        for name, details in order.get_items().items():
            quantity = details["quantity"]
            price = details["price"]

            lines.append(
                f"{name} | تعداد: {quantity} | "
                f"قیمت واحد: {price:,} | "
                f"جمع: {quantity * price:,} تومان"
            )

        lines.extend([
            "-" * 40,
            f"مبلغ کل: {order.calculate_total():,} تومان"
        ])

        invoice_text.insert("1.0", "\n".join(lines))
        invoice_text.tag_configure("right", justify="right")
        invoice_text.tag_add("right", "1.0", "end")
        invoice_text.config(state="disabled")
        def save_invoice():
            try:
                file_path = filedialog.asksaveasfilename(
                    parent=invoice_window,
                    title="ذخیره فاکتور",
                    defaultextension=".txt",
                    initialfile=f"invoice_{order.order_id}.txt",
                    filetypes=[
                        ("Text files", "*.txt")
                    ]
                )

                if not file_path:
                    return

                with open(file_path, "w", encoding="utf-8-sig") as file:
                    file.write("\n".join(lines))

                messagebox.showinfo(
                    "ذخیره موفق",
                    "فاکتور با موفقیت ذخیره شد.",
                    parent=invoice_window
                )

            except OSError as error:
                messagebox.showerror(
                    "خطای ذخیره فایل",
                    str(error),
                    parent=invoice_window
                )

        save_button = ttk.Button(
            invoice_window,
            text="ذخیره فاکتور",
            command=save_invoice
        )

        save_button.pack(
            side="bottom",
            pady=10
        )

        invoice_text.pack(
            side="top",
            fill="both",
            expand=True,
            padx=10,
            pady=5
        )

    finalize_button = ttk.Button(
        order_tab,
        text="ثبت نهایی سفارش",
        command=finalize_order
    )

    finalize_button.pack(
        pady=5,
        before=order_table
    )

    def show_customer_details():

        selected = customers_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب مشتری",
                "لطفاً ابتدا یک مشتری از جدول انتخاب کنید."
            )
            return

        values = customers_table.item(selected[0], "values")
        customer_id = int(values[0])

        try:
            customer = Customer.load_customer(customer_id)

            details_window = tk.Toplevel(root)
            details_window.title("جزئیات مشتری")
            details_window.geometry("450x350")
            details_window.transient(root)

            details = [
                ("شناسه مشتری", customer.customer_id),
                ("نام", customer.first_name),
                ("نام خانوادگی", customer.last_name),
                ("شماره تماس", customer.phone_number),
                ("آدرس", customer.address),
                ("جنسیت", customer.gender)
            ]

            for row_index, (label_text, value) in enumerate(details):

                ttk.Label(
                    details_window,
                    text=f"{label_text}:",
                    font=("Tahoma", 11, "bold")
                ).grid(
                    row=row_index,
                    column=1,
                    padx=15,
                    pady=12,
                    sticky="e"
                )

                ttk.Label(
                    details_window,
                    text=str(value),
                    font=("Tahoma", 11),
                    wraplength=240
                ).grid(
                    row=row_index,
                    column=0,
                    padx=15,
                    pady=12,
                    sticky="e"
                )

        except (ValueError, OSError, EOFError, pickle.UnpicklingError) as error:
            messagebox.showerror(
                "خطا در نمایش اطلاعات",
                str(error)
            )

    details_button = ttk.Button(
        admin_tab,
        text="مشاهده جزئیات مشتری",
        command=show_customer_details
    )

    details_button.pack(pady=5)

    def edit_customer():
        selected = customers_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب مشتری",
                "لطفاً ابتدا یک مشتری را انتخاب کنید."
            )
            return

        values = customers_table.item(selected[0], "values")
        customer_id = int(values[0])

        try:
            customer = Customer.load_customer(customer_id)

        except (OSError, ValueError, EOFError, pickle.UnpicklingError) as error:
            messagebox.showerror("خطا", str(error))
            return

        edit_window = tk.Toplevel(root)
        edit_window.title("ویرایش اطلاعات مشتری")
        edit_window.geometry("500x420")
        edit_window.transient(root)

        fields = [
            ("نام", customer.first_name),
            ("نام خانوادگی", customer.last_name),
            ("شماره تماس", customer.phone_number),
            ("آدرس", customer.address)
        ]

        entries = {}

        for row_index, (label_text, value) in enumerate(fields):
            ttk.Label(
                edit_window,
                text=label_text
            ).grid(
                row=row_index,
                column=1,
                padx=15,
                pady=12
            )

            entry = ttk.Entry(edit_window, width=35)
            entry.insert(0, value)
            entry.grid(
                row=row_index,
                column=0,
                padx=15,
                pady=12
            )

            entries[label_text] = entry

        ttk.Label(
            edit_window,
            text="جنسیت"
        ).grid(row=4, column=1, padx=15, pady=12)

        gender_var = tk.StringVar(value=customer.gender)

        gender_box = ttk.Combobox(
            edit_window,
            textvariable=gender_var,
            values=["مرد", "زن", "سایر"],
            state="readonly",
            width=32
        )
        gender_box.grid(row=4, column=0, padx=15, pady=12)

        def save_changes():
            try:
                updated_customer = Customer(
                    customer_id=customer.customer_id,
                    first_name=entries["نام"].get().strip(),
                    last_name=entries["نام خانوادگی"].get().strip(),
                    phone_number=entries["شماره تماس"].get().strip(),
                    address=entries["آدرس"].get().strip(),
                    gender=gender_var.get()
                )

                updated_customer.update_customer()

                refresh_customers()

                messagebox.showinfo(
                    "ویرایش موفق",
                    "اطلاعات مشتری با موفقیت به‌روزرسانی شد.",
                    parent=edit_window
                )

                edit_window.destroy()

            except (ValueError, TypeError, OSError, EOFError,
                    pickle.UnpicklingError) as error:
                messagebox.showerror(
                    "خطای ویرایش",
                    str(error),
                    parent=edit_window
                )

        save_button = ttk.Button(
            edit_window,
            text="ذخیره تغییرات",
            command=save_changes
        )

        save_button.grid(
            row=5,
            column=0,
            columnspan=2,
            pady=20
        )

    edit_button = ttk.Button(
        admin_tab,
        text="ویرایش مشتری",
        command=edit_customer
    )

    edit_button.pack(pady=5)

    def delete_selected_customer():
        nonlocal current_customer

        selected = customers_table.selection()

        if not selected:
            messagebox.showwarning(
                "انتخاب مشتری",
                "لطفاً ابتدا یک مشتری را انتخاب کنید."
            )
            return

        values = customers_table.item(selected[0], "values")
        customer_id = int(values[0])

        try:
            customer = Customer.load_customer(customer_id)

            # بررسی سفارش‌های قبلی
            import os

            orders_file = "data/orders.pkl"

            if os.path.isfile(orders_file):
                with open(orders_file, "rb") as file:
                    orders = pickle.load(file)

                if not isinstance(orders, dict):
                    raise ValueError(
                        "ساختار فایل سفارش‌ها معتبر نیست."
                    )

                for order in orders.values():
                    if order.customer.customer_id == customer_id:
                        messagebox.showwarning(
                            "حذف غیرمجاز",
                            "این مشتری سفارش ثبت‌شده دارد و "
                            "برای حفظ سوابق امکان حذف او وجود ندارد."
                        )
                        return

            # دریافت تأیید کاربر
            confirmed = messagebox.askyesno(
                "تأیید حذف مشتری",
                f"آیا از حذف مشتری "
                f"{customer.first_name} {customer.last_name} "
                f"با شناسه {customer_id} مطمئن هستید؟"
            )

            if not confirmed:
                return

            Customer.delete_customer(customer_id)

            if (
                current_customer is not None
                and current_customer.customer_id == customer_id
            ):
                current_customer = None

            refresh_customers()

            messagebox.showinfo(
                "حذف موفق",
                "مشتری با موفقیت حذف شد."
            )

        except (
            OSError,
            ValueError,
            TypeError,
            EOFError,
            pickle.UnpicklingError
        ) as error:
            messagebox.showerror(
                "خطای حذف مشتری",
                str(error)
            )

    delete_button = ttk.Button(
        admin_tab,
        text="حذف مشتری",
        command=delete_selected_customer
    )

    delete_button.pack(pady=5)


    root.mainloop()


if __name__ == "__main__":
    run_app()